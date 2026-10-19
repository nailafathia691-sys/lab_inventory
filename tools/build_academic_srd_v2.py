import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from academic_docx_builder import (
    setup_academic_page, add_academic_cover, add_kata_pengantar,
    add_bab_heading, add_subbab_heading, add_body_paragraph,
    add_bullet_item, add_academic_image, add_academic_table,
    add_daftar_pustaka
)

IMG_DEV = os.path.abspath("docs/evidence/device")
IMG_EVID = os.path.abspath("docs/evidence")

DAFTAR_PUSTAKA_SRD = [
    "Budhayanti, C. I. S. (2023). Penggunaan Mobile Learning dan Cloud Database pada Perkuliahan Pendidikan Komputer. Jurnal Pendidikan Terbuka dan Jarak Jauh, 24(1), 45-56.",
    "Kemendikbudristek. (2022). Panduan Pembelajaran dan Asesmen Pendidikan Vokasi dan Komputer. Jakarta: BSKAP Kemendikbudristek.",
    "PostgreSQL Global Development Group. (2023). PostgreSQL 15 Documentation: Row Level Security Policies and Concurrency Control. PostgreSQL Official Documentation Hub.",
    "Sommerville, I. (2019). Software Engineering (10th ed.). London: Pearson Education.",
    "Supabase Inc. (2024). Supabase Documentation: Auth, Realtime Database, and Stored Procedures via PostgreSQL RPC. San Francisco: Supabase Docs Hub.",
    "Team Flutter. (2024). Flutter Documentation: Material 3 Design Specifications and State Management with IndexedStack. Mountain View: Google Open Source.",
    "Wulandari, L. P. A. O., & Agustini, K. (2024). Pengembangan Sistem Manajemen Laboratorium Berbasis Mobile Menggunakan Arsitektur Clean Architecture. Jurnal Pendidikan Teknologi dan Kejuruan, 23(1), 89-101."
]

def build_academic_srd_v2():
    doc = docx.Document()
    setup_academic_page(doc)

    # Cover
    add_academic_cover(
        doc,
        jenis_dokumen="SOFTWARE REQUIREMENT DOCUMENT (SRD) — VERSION 2.0",
        judul_utama="SISTEM INFORMASI MANAJEMEN INVENTARIS DAN PEMINJAMAN ALAT LABORATORIUM KOMPUTER BERBASIS MOBILE (LABINVENTORY)",
        sub_judul="Spesifikasi Kebutuhan Perangkat Lunak Terintegrasi Cloud Supabase PostgreSQL Realtime dan Row Level Security (RLS)",
        tgl_tahun="2026"
    )

    # Kata Pengantar
    add_kata_pengantar(
        doc,
        judul_dokumen="Software Requirement Document (SRD) Version 2.0: Sistem Informasi Manajemen Inventaris dan Peminjaman Alat Laboratorium Komputer Berbasis Mobile (LabInventory)",
        tgl_str="19 Oktober 2026"
    )

    # Riwayat Revisi
    h_rev = doc.add_paragraph()
    h_rev.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_rev.paragraph_format.space_before = Pt(12)
    h_rev.paragraph_format.space_after = Pt(14)
    r_rev = h_rev.add_run("RIWAYAT REVISI DOKUMEN")
    r_rev.font.name = "Times New Roman"
    r_rev.font.size = Pt(12)
    r_rev.bold = True

    add_body_paragraph(
        doc,
        "Setiap modifikasi arsitektural dan spesifikasi kebutuhan dicatat secara berkala guna menjamin transparansi serta "
        "kepatuhan implementasi terhadap rancangan teknis yang disepakati."
    )

    rev_headers = ["Versi", "Tanggal", "Deskripsi Perubahan Spesifikasi", "Penulis"]
    rev_data = [
        ["1.0", "03 Sept 2026", "Draf awal dokumen SRD: Pendahuluan, arsitektur dasar, rancangan use case, spesifikasi antarmuka awal, dan konsep awal MySQL.", "Naila Fathia"],
        ["2.0", "19 Okt 2026", "Pembaruan Menyeluruh (v2.0): Migrasi arsitektur dari PHP/MySQL ke Flutter + Supabase (PostgreSQL 15, Supabase Auth, Row Level Security / RLS, dan RPC Transactional Function untuk integritas stok atomik), skema UUID, penghapusan dummy data, dan penyertaan hasil pengujian perangkat fisik Android.", "Naila Fathia"]
    ]
    add_academic_table(doc, "Tabel Riwayat Revisi Dokumen Spesifikasi Kebutuhan", rev_headers, rev_data, col_widths=[0.8, 1.2, 3.5, 1.0])
    doc.add_page_break()

    # BAB I
    add_bab_heading(doc, "I", "PENDAHULUAN")
    add_subbab_heading(doc, "1.1", "Latar Belakang Sistem")
    add_body_paragraph(
        doc,
        "Laboratorium Komputer pada Program Studi Pendidikan Teknologi Informasi (PTI), FKIP Universitas Lampung, "
        "mengelola puluhan perangkat praktikum seperti laptop mahasiswa, proyektor LCD, switch jaringan, instrumen kabel, "
        "serta perangkat multimedia. Sebelumnya, proses pendataan barang dan permohonan peminjaman praktikum masih mengandalkan "
        "buku catatan fisik formulir kertas."
    )
    add_body_paragraph(
        doc,
        "Metode manual tersebut memiliki kelemahan mendasar: tingginya risiko kehilangan catatan, kesulitan melacak status ketersediaan "
        "stok alat secara instan, tidak adanya verifikasi kondisi saat pengembalian, serta antrean administratif yang memakan waktu mahasiswa "
        "dan laboran. Oleh karena itu, dikembangkan sistem aplikasi mobile 'LabInventory' guna memodernisasi tata kelola inventaris "
        "secara digital, transparan, dan real-time."
    )

    add_subbab_heading(doc, "1.2", "Gambaran Umum Sistem")
    add_body_paragraph(
        doc,
        "LabInventory adalah aplikasi mobile berbasis Flutter yang terintegrasi secara langsung dengan PostgreSQL Cloud Database via Supabase. "
        "Sistem melayani dua peran utama:"
    )
    add_bullet_item(doc, "Mengecek ketersediaan stok live, melihat spesifikasi alat, mengajukan permohonan peminjaman praktikum, memantau alur persetujuan, dan melihat tenggat waktu pengembalian.", bold_prefix="Mahasiswa: ")
    add_bullet_item(doc, "Memverifikasi pengajuan peminjaman (approve/reject), mengelola stok dan kondisi fisik inventaris, memproses pengembalian alat (return), serta meninjau riwayat transaksi seluruh lab.", bold_prefix="Laboran / Admin: ")

    add_subbab_heading(doc, "1.3", "Tujuan dan Manfaat Sistem")
    add_body_paragraph(
        doc,
        "Tujuan utama pengembangan sistem meliputi:"
    )
    add_bullet_item(doc, "Mentransformasi proses peminjaman manual menjadi sistem digital terpusat yang aman dan akuntabel.")
    add_bullet_item(doc, "Mencegah kesalahan alokasi unit (double-booking) melalui transaksi database yang atomik (ACID via RPC).")
    add_bullet_item(doc, "Menyediakan audit trail transparan mengenai riwayat keluar-masuk barang serta kondisi fisik alat pasca praktikum.")

    # BAB II
    add_bab_heading(doc, "II", "ARSITEKTUR DAN LINGKUNGAN OPERASIONAL")
    add_subbab_heading(doc, "2.1", "Arsitektur Sistem & Technology Stack")
    add_body_paragraph(
        doc,
        "Sistem menerapkan arsitektur Clean Architecture pada sisi klien Flutter dengan pola Repository Pattern yang berkomunikasi aman "
        "ke Supabase Cloud Platform. Keamanan data tidak diletakkan pada logika klien, melainkan ditegakkan langsung di level database engine "
        "menggunakan Row Level Security (RLS) PostgreSQL."
    )

    arch_img = os.path.join(IMG_EVID, "diagram_arsitektur_sistem.png")
    if os.path.exists(arch_img):
        add_academic_image(doc, arch_img, "Gambar 2.1: Diagram Arsitektur Cloud, Gateway Keamanan, dan Database Supabase PostgreSQL 15", width=Inches(5.4))

    tech_headers = ["Lapisan / Layer", "Teknologi & Versi", "Fungsi & Tanggung Jawab"]
    tech_data = [
        ["Mobile Client", "Flutter 3.13+ / Dart 3+", "Antarmuka responsif Material Design 3, state management, offline-fallback demo."],
        ["Autentikasi", "Supabase Auth (GoTrue)", "Manajemen sesi aman berbasis token JWT, sanitasi pendaftaran, login role-based."],
        ["Cloud Database", "PostgreSQL 15 (Supabase Cloud)", "Penyimpanan data relasional terpusat, foreign keys, index performa."],
        ["Otorisasi & Keamanan", "Row Level Security (RLS)", "Aturan hak akses baris: Mahasiswa hanya membaca miliknya, Admin mengelola seluruh lab."],
        ["Integritas Transaksi", "PostgreSQL RPC Functions", "Stored procedure atomik (`approve_borrowing`, `return_borrowing`) bergaransi ACID."],
        ["Aset Grafis & Media", "Vector Mipmaps & Cloud Media", "Penyajian gambar foto alat dan logo resolusi tinggi tanpa distorsi layout."]
    ]
    add_academic_table(doc, "Tabel 2.1: Spesifikasi Technology Stack Sistem LabInventory", tech_headers, tech_data, col_widths=[1.5, 2.0, 3.0])

    # BAB III
    add_bab_heading(doc, "III", "SPESIFIKASI PENGGUNA DAN USE CASE")
    add_subbab_heading(doc, "3.1", "Model Diagram Use Case")
    add_body_paragraph(
        doc,
        "Diagram Use Case di bawah ini mengilustrasikan pembagian hak akses fitur dan batasan operasional sistem antara Mahasiswa dan Laboran:"
    )

    usecase_img = os.path.join(IMG_EVID, "diagram_usecase_sistem.png")
    if os.path.exists(usecase_img):
        add_academic_image(doc, usecase_img, "Gambar 3.1: Diagram Use Case & Pembagian Peran Aplikasi Mobile LabInventory", width=Inches(5.4))

    uc_headers = ["Kode", "Nama Use Case", "Aktor Utama", "Deskripsi Alur Sistem"]
    uc_data = [
        ["UC-01", "Autentikasi & Validasi Peran", "Mahasiswa, Laboran", "Pengguna memasukkan email dan kata sandi. Supabase Auth memvalidasi token dan server mencocokkan role pada tabel `profiles`."],
        ["UC-02", "Pencarian & Filter Katalog", "Mahasiswa, Laboran", "Melihat daftar alat aktif, filter per kategori (Komputer, Elektronik, Aksesoris), dan pencarian real-time."],
        ["UC-03", "Pengajuan Peminjaman Alat", "Mahasiswa", "Mahasiswa memilih alat, menentukan jumlah unit, tanggal pinjam & kembali, serta menginput tujuan praktikum."],
        ["UC-04", "Verifikasi & Persetujuan (Approve/Reject)", "Laboran", "Laboran meninjau pengajuan. Jika disetujui, RPC memverifikasi stok dan mengalokasikan unit secara atomik."],
        ["UC-05", "Pencatatan Pengembalian & Cek Fisik", "Laboran", "Laboran menginspeksi fisik alat (Kondisi Baik atau Rusak), mencatat evaluasi, dan sistem mengembalikan stok siap pakai / stok rusak."],
        ["UC-06", "Pendaftaran Inventaris Baru", "Laboran", "Laboran menambahkan alat baru dengan kode inventaris otomatis (`INV-[KAT]-[RANDOM]`), kondisi, stok awal, dan lokasi rak."],
        ["UC-07", "Audit Log & Riwayat Transaksi", "Laboran, Mahasiswa", "Mahasiswa melihat riwayat transaksi pribadinya, sedangkan Laboran melihat audit log seluruh transaksi laboratorium."]
    ]
    add_academic_table(doc, "Tabel 3.1: Matriks Use Case Utama Sistem LabInventory", uc_headers, uc_data, col_widths=[0.8, 1.8, 1.2, 2.7])

    # BAB IV
    add_bab_heading(doc, "IV", "KEBUTUHAN FUNGSIONAL & NON-FUNGSIONAL")
    add_subbab_heading(doc, "4.1", "Kebutuhan Fungsional (Functional Requirements)")
    fr_headers = ["Kode", "Nama Kebutuhan", "Spesifikasi Implementasi Sistem"]
    fr_data = [
        ["FR-01", "Manajemen Sesi & Role Guard", "Sistem membatasi navigasi mahasiswa hanya pada fitur peminjam dan admin pada fitur manajemen lab."],
        ["FR-02", "Monitoring Stok Real-time", "Stok inventaris terupdate secara live di cloud saat transaksi disetujui atau dikembalikan."],
        ["FR-03", "Validasi Tanggal Kalender", "Tanggal pinjam minimal hari ini (WIB) dan tanggal kembali tidak boleh lebih awal dari tanggal pinjam."],
        ["FR-04", "Pencegahan Over-borrowing", "Pengajuan tidak dapat melebihi batas stok tersedia pada saat form dikirim maupun saat diapprove."],
        ["FR-05", "Klasifikasi Pengembalian Rusak", "Unit yang dikembalikan dalam kondisi rusak otomatis dialihkan ke `stok_rusak` tanpa mengubah status barang sehat lainnya."],
        ["FR-06", "Pemberian Kode Otomatis", "Pendaftaran alat baru menghasilkan kode unik format standar laboratorium."],
        ["FR-07", "Dashboard Metrik Laboran", "Menghitung total aset aktif, jumlah antrean pending, dan rekapitulasi status secara instan."]
    ]
    add_academic_table(doc, "Tabel 4.1: Spesifikasi Kebutuhan Fungsional Sistem (FR)", fr_headers, fr_data, col_widths=[0.8, 1.8, 3.9])

    add_subbab_heading(doc, "4.2", "Kebutuhan Non-Fungsional (Non-Functional Requirements)")
    add_bullet_item(doc, "Seluruh pertukaran data dilindungi protokol TLS 1.3. Akses tabel dibatasi kebijakan RLS PostgreSQL. Password di-hash menggunakan algoritma Bcrypt via GoTrue.", bold_prefix="NFR-01 Keamanan: ")
    add_bullet_item(doc, "Eksekusi persetujuan dan pengembalian dijalankan dalam satu blok transaksi database PostgreSQL (`BEGIN ... COMMIT`) untuk mencegah anomali data.", bold_prefix="NFR-02 Integritas Transaksi: ")
    add_bullet_item(doc, "Setiap request penulisan menyertakan UUID unik sebagai idempotency key untuk mencegah pengiriman ganda akibat sinyal seluler yang tidak stabil.", bold_prefix="NFR-03 Idempotensi Jaringan: ")
    add_bullet_item(doc, "Antarmuka dirancang adaptif terhadap berbagai ukuran layar smartphone Android, mendukung Safe Area (bebas tumpang tindih keyboard/notch).", bold_prefix="NFR-04 Responsivitas Antarmuka: ")
    add_bullet_item(doc, "Waktu respons server untuk query inventaris rata-rata di bawah 500ms pada koneksi jaringan 4G stabil.", bold_prefix="NFR-05 Performa & Kecepatan: ")

    # BAB V
    add_bab_heading(doc, "V", "PERANCANGAN STRUKTUR BASIS DATA")
    add_subbab_heading(doc, "5.1", "Tabel profiles (Profil Pengguna)")
    db_p_headers = ["Nama Kolom", "Tipe Data", "Batasan / Constraint", "Keterangan"]
    db_p_data = [
        ["id", "UUID", "PRIMARY KEY, REFERENCES auth.users(id)", "ID unik akun pengguna dari Supabase Auth"],
        ["nama", "TEXT", "NOT NULL", "Nama lengkap mahasiswa atau laboran"],
        ["npm", "TEXT", "UNIQUE, NULLABLE", "Nomor Pokok Mahasiswa (khusus mahasiswa)"],
        ["role", "TEXT", "NOT NULL, CHECK (role IN ('student','admin'))", "Hak akses peran pengguna dalam sistem"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Waktu pembuatan akun"]
    ]
    add_academic_table(doc, "Tabel 5.1: Struktur Tabel Basis Data profiles", db_p_headers, db_p_data, col_widths=[1.2, 1.2, 2.0, 2.1])

    add_subbab_heading(doc, "5.2", "Tabel items (Inventaris Alat Laboratorium)")
    db_i_headers = ["Nama Kolom", "Tipe Data", "Batasan / Constraint", "Keterangan"]
    db_i_data = [
        ["id", "UUID", "PRIMARY KEY, DEFAULT gen_random_uuid()", "ID unik barang inventaris"],
        ["nama_barang", "TEXT", "NOT NULL", "Nama alat (cth: Laptop Lenovo, Proyektor Epson)"],
        ["kode_barang", "TEXT", "NOT NULL, UNIQUE", "Kode inventaris lab (cth: INV-CMP-0008)"],
        ["kategori", "TEXT", "NOT NULL", "Komputer, Elektronik, Aksesoris, Audio"],
        ["kondisi", "TEXT", "NOT NULL, DEFAULT 'Baik'", "Sangat Baik, Baik, Perlu Perbaikan"],
        ["stok", "INTEGER", "NOT NULL, CHECK (stok >= 0)", "Jumlah unit siap pakai yang tersedia"],
        ["stok_rusak", "INTEGER", "NOT NULL, DEFAULT 0", "Jumlah unit yang rusak dan dalam perbaikan"],
        ["lokasi", "TEXT", "NOT NULL", "Lokasi rak / lemari di laboratorium PTI"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Waktu pencatatan aset"]
    ]
    add_academic_table(doc, "Tabel 5.2: Struktur Tabel Basis Data items", db_i_headers, db_i_data, col_widths=[1.2, 1.2, 2.0, 2.1])

    add_subbab_heading(doc, "5.3", "Tabel borrowings (Transaksi Peminjaman)")
    db_b_headers = ["Nama Kolom", "Tipe Data", "Batasan / Constraint", "Keterangan"]
    db_b_data = [
        ["id", "UUID", "PRIMARY KEY, DEFAULT gen_random_uuid()", "ID unik transaksi peminjaman"],
        ["user_id", "UUID", "NOT NULL, REFERENCES profiles(id)", "ID mahasiswa yang mengajukan peminjaman"],
        ["item_id", "UUID", "NOT NULL, REFERENCES items(id)", "ID alat yang dipinjam"],
        ["nama_peminjam", "TEXT", "NOT NULL", "Nama lengkap peminjam"],
        ["nama_barang", "TEXT", "NOT NULL", "Nama barang yang dipinjam"],
        ["jumlah", "INTEGER", "NOT NULL, CHECK (jumlah > 0)", "Jumlah unit yang diajukan"],
        ["tgl_pinjam", "TIMESTAMPTZ", "NOT NULL", "Tanggal awal masa peminjaman"],
        ["tgl_kembali", "TIMESTAMPTZ", "NOT NULL", "Tenggat waktu pengembalian alat"],
        ["keperluan", "TEXT", "NOT NULL", "Alasan praktikum / kegiatan akademik"],
        ["status", "TEXT", "CHECK (status IN ('pending','approved','rejected','returned'))", "Status persetujuan dan transaksi"],
        ["catatan_kondisi", "TEXT", "NULLABLE", "Catatan hasil inspeksi fisik oleh laboran"],
        ["approved_at", "TIMESTAMPTZ", "NULLABLE", "Waktu persetujuan oleh laboran"],
        ["returned_at", "TIMESTAMPTZ", "NULLABLE", "Waktu pengembalian aktual barang"],
        ["created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Waktu pengajuan dikirim"]
    ]
    add_academic_table(doc, "Tabel 5.3: Struktur Tabel Basis Data borrowings", db_b_headers, db_b_data, col_widths=[1.3, 1.1, 2.0, 2.1])

    # BAB VI
    add_bab_heading(doc, "VI", "SPESIFIKASI ANTARMUKA PENGGUNA")
    add_body_paragraph(
        doc,
        "Berikut merupakan dokumentasi antarmuka pengguna hasil kompilasi langsung pada perangkat fisik Android (Physical Device Testing):"
    )

    screenshots = [
        ("01-login.png", "Gambar 6.1: Antarmuka Autentikasi Pengguna (Login Mahasiswa & Laboran dengan Opsi Demo)"),
        ("10-catalog-search.png", "Gambar 6.2: Katalog Inventaris Real-time Dilengkapi Fitur Pencarian & Filter Kategori"),
        ("08-laptop-detail.png", "Gambar 6.3: Halaman Rincian Alat (Spesifikasi, Lokasi Lemari, dan Ketersediaan Stok)"),
        ("02-student-form.png", "Gambar 6.4: Formulir Digital Pengajuan Peminjaman Alat Mahasiswa"),
        ("03-admin-approved.png", "Gambar 6.5: Dashboard Laboran dengan Fitur Verifikasi & Persetujuan Cepat"),
        ("06-damaged-return.png", "Gambar 6.6: Dialog Verifikasi Pengembalian Alat dengan Inspeksi Kondisi Fisik"),
        ("09-admin-created-item.png", "Gambar 6.7: Formulir Pendaftaran Alat Baru oleh Laboran"),
        ("11-final-history.png", "Gambar 6.8: Audit Log & Riwayat Lengkap Seluruh Transaksi Laboratorium")
    ]

    for filename, caption in screenshots:
        full_path = os.path.join(IMG_DEV, filename)
        if os.path.exists(full_path):
            add_academic_image(doc, full_path, caption, width=Inches(3.1))

    # BAB VII
    add_bab_heading(doc, "VII", "PENUTUP")
    add_subbab_heading(doc, "7.1", "Kesimpulan Keterlaksanaan Sistem")
    add_body_paragraph(
        doc,
        "Pembaruan SRD versi 2.0 ini berhasil menyelaraskan dokumen perancangan dengan implementasi final aplikasi LabInventory. "
        "Penggunaan Flutter dan Supabase PostgreSQL terbukti mampu menyelesaikan permasalahan pencatatan manual di Laboratorium "
        "Komputer PTI FKIP Unila secara aman, real-time, dan mudah dioperasikan baik oleh mahasiswa maupun laboran."
    )

    add_daftar_pustaka(doc, DAFTAR_PUSTAKA_SRD)

    target_v2 = os.path.abspath("SOFTWARE REQUIREMENT DOCUMENT_v2.docx")
    target_root = os.path.abspath("SOFTWARE REQUIREMENT DOCUMENT.docx")
    doc.save(target_v2)
    doc.save(target_root)
    print("SRD v2 generated matching Unila academic specifications successfully!")

if __name__ == "__main__":
    build_academic_srd_v2()
